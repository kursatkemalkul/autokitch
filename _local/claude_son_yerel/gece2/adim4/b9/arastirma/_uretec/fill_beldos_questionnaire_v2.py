"""Create the completed Beldos research questionnaire and the sample-video sheet.

The supplier's questionnaire is retained as the source document. Engineering
targets and proposed sample specs are never presented as measured properties.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1] / "TEDARIK"
SOURCE = ROOT / "Beldos_Questionnaire_Original.docx"
SUPPLIER = ROOT / "Beldos_Questionnaire_AUTOKITCH_RESEARCH_v2.docx"
VIDEO = ROOT / "AUTOKITCH_Beldos_Numune_Video_Formu_v1.docx"


def put(cell, value, size=8.1):
    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(size)


def supplier_form():
    doc = Document(SOURCE)
    hdr = doc.sections[0].header.tables[0]
    put(hdr.cell(0, 1), "AUTOKITCH (research project)", 9)
    put(hdr.cell(0, 3), "Kemal Kul | k.kemalkul@gmail.com", 9)
    t = doc.tables[0]
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

    rows = [
        (
            "Pizza sauce",
            "80 g / deposit*",
            "2 Semi-liquid*; spoon-flow video pending",
            "4 °C planned*; actual sample °C pending",
            "Crushed tomato/small flecks; maximum piece not measured",
            "Mixing bowl not selected; hopper target 10–15 kg*",
            "0 actual; planned batches/day not fixed",
            "0 actual; if full hopper = batch, ~125–188 portions*",
        ),
        (
            "Lahmacun mixture",
            "110 g / deposit*; published example: 100 g [1]",
            "3 Soft paste* with particles; video pending",
            "4 °C planned*; actual sample °C pending",
            "Raw ground beef + onion + pepper paste; largest piece not measured [1]",
            "Mixing bowl not selected; 2 separate hopper targets, 21.6 kg each*",
            "0 actual; planned batches/day not fixed",
            "0 actual; if full hopper = batch, ~196 portions / hopper*",
        ),
        (
            "Diced raw beef",
            "145 g / deposit*",
            "Discrete chunks; Beldos liquid/paste scale does not apply",
            "2 °C planned*; actual sample °C pending",
            "Wet raw beef; 15–20 mm requested trial cut* (not ordinary 30–40 mm diced beef [2])",
            "No mixing bowl; hopper target 5.8 kg*",
            "0 actual; planned batches/day not fixed",
            "0 actual; if full hopper = batch, 40 portions*",
        ),
        (
            "Raw minced beef",
            "160 g / deposit*",
            "4 Heavy paste* / granular mass; video pending",
            "2 °C planned*; actual sample °C pending",
            "Raw ground beef; 3 mm grinder plate proposed*; actual piece size unmeasured",
            "No mixing bowl; hopper target 6.4 kg*",
            "0 actual; planned batches/day not fixed",
            "0 actual; if full hopper = batch, 40 portions*",
        ),
    ]
    for i, values in enumerate(rows, start=1):
        row = t.rows[i]
        row.height = None
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for j, value in enumerate(values):
            put(row.cells[j], value, 7.7)

    # The original form has five product rows; the inquiry is for exactly four.
    fifth = t.rows[5]
    fifth.height = None
    fifth.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = fifth.cells[0].merge(fifth.cells[7])
    put(cell, "No fifth product in this inquiry — four independent product paths requested.", 8)

    put(t.cell(6, 3), "Not applicable — no potato purée.", 9)
    put(t.cell(7, 3), "0 actual production days/week. Planned schedule not decided.", 9)
    put(t.cell(8, 3), "No operating baseline or growth forecast. Engineering capacity target: ~120 finished products/hour in total*.", 9)
    put(t.cell(9, 3), "Onto an uncooked round pide/lahmacun dough base, Ø280 mm design target*. Sketch and fresh-product videos to follow.", 9)
    put(t.cell(10, 3), "No for baseline: four separate product hoppers, paths and outlets. Supplier alternatives welcome.", 9)
    put(t.cell(11, 3), "Internet search.", 9)

    para = t.cell(12, 0).paragraphs[-1]
    for line in (
        "\nAUTOKITCH: * VARSAYIM / design or trial target, NOT a measured product property. The line is not operating. Current manual-process video does not exist; fresh sample spoon-flow videos are being prepared. Do not select equipment from this sheet alone.",
        "\n[1] Aydemir et al., Food and Health 11(2), 2025, DOI 10.3153/FH25011 (published example recipe/100 g filling; not our recipe). [2] Turkish Meat and Milk Board: esk.gov.tr/tr/10998/ET (ordinary diced beef 3–4 cm).",
    ):
        run = para.add_run(line)
        run.font.size = Pt(7.3)
        run.bold = line.startswith("\nAUTOKITCH")
    doc.save(SUPPLIER)


def text(doc, value, style=None, boldlead=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    if boldlead and value.startswith(boldlead):
        p.add_run(boldlead).bold = True
        p.add_run(value[len(boldlead):])
    else:
        p.add_run(value)
    return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", fill)
    tcPr.append(shd)


def mini_table(doc, data, widths=None):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Cm(4)
    table.columns[1].width = Cm(12.7)
    for n, (key, val) in enumerate(data):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        row.cells[0].width = Cm(4)
        row.cells[1].width = Cm(12.7)
        put(row.cells[0], key, 9)
        put(row.cells[1], val, 9)
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True
        if n % 2 == 0:
            shade(row.cells[0], "EFF3F6")
            shade(row.cells[1], "EFF3F6")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def header(doc, title, subtitle):
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(18, 57, 73)
    p.paragraph_format.space_after = Pt(3)
    p2 = doc.add_paragraph()
    p2.add_run(subtitle)
    p2.paragraph_format.space_after = Pt(12)


def product(doc, name, target, source, sample, steps, fields):
    doc.add_heading(name, level=1)
    mini_table(doc, [
        ("Doz hedefi", target),
        ("Örnek ürün", sample),
        ("Dayanak", source),
    ])
    for step in steps:
        text(doc, "□ " + step)
    mini_table(doc, fields)


def video_form():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.55)
    sec.left_margin = sec.right_margin = Cm(2.1)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor(32, 45, 55)
    for nm in ("Title", "Heading 1", "Heading 2"):
        styles[nm].font.name = "Aptos Display"
        styles[nm].font.color.rgb = RGBColor(18, 57, 73)
    styles["Title"].font.size = Pt(18)
    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 1"].paragraph_format.space_before = Pt(10)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)

    header(doc, "AUTOKITCH | NUMUNE VE VİDEO FORMU", "Beldos'a gönderilecek dört ürün örneğinin hazırlanması ve kaydı · v1 · 25.09.2026")
    text(doc, "Bu form üretim reçetesi değildir. Mevcut makine yalnızca araştırma/tasarım aşamasındadır. Aşağıdaki gramajlar AUTOKITCH tasarım hedefidir; ürünün gerçek akışkanlığı ve sıcaklığı videoda ölçülüp kaydedilecektir.")
    doc.add_heading("ORTAK ÇEKİM DÜZENİ", level=1)
    for s in (
        "Her ürün için ayrı, temiz kap ve ayrı video kullan. İlk karede ürün adı, marka/tedarikçi, tarih ve yaklaşık 1 kg numuneyi göster.",
        "Dijital teraziyi ve termometreyi kadraja al; ürünün gerçek sıcaklığını yaz. Özellikle çiğ et/harcı çekim boyunca soğuk tut; bekletip ısıtarak akışkanlık gösterme.",
        "Büyük servis kaşığını ürüne daldır, kaldır, doğal akışını kesmeden 10–15 saniye göster. Beldos küçük kaşık istemiyor. Aynı hareketi 2–3 kez tekrarla.",
        "Ölçek için cetvel göster; iri parçaları ayrı yakın çek. Sıvı ayrışması/et suyu varsa kabın dibini de göster, karıştırmadan önce ve sonra kısa çekim yap.",
        "Hedef porsiyonu tart ve çiğ Ø280 mm yuvarlak hamur üzerinde nasıl yayıldığını göster (hamur ölçüsü VARSAYIM/tasarım hedefi). Gerçek elle serme varsa bunu da kaydet.",
        "Dört videoyu ürün adıyla kaydet; dosya adına tarih ve parti numarası ekle. Her video ile aşağıdaki gerçek ölçüm satırlarını birlikte gönder.",
    ):
        text(doc, "□ " + s)
    doc.add_heading("ORTAK KAYIT", level=1)
    mini_table(doc, [
        ("Tarih / hazırlayan", "___________________________ / ___________________________"),
        ("Hamur çapı", "Gerçekte ölçülen ______ mm (tasarım hedefi Ø280 mm)") ,
        ("Kayıt yöntemi", "Marka/ürün etiketi + terazi + termometre + büyük kaşık + cetvel"),
        ("Eksik olabilecek veri", "Gerçek viskozite (Pa·s), en iri parça ve et suyu ayrışması: numunede görülene kadar bilinmiyor."),
    ])

    doc.add_page_break()
    product(doc, "1 | PİZZA SOSU", "80 g / pide · VARSAYIM: tasarım dozu", "Üreticiye ait başka bir ticari pizza sosu spesifikasyonu: 7–11 °Bx, Bostwick 4–8 cm/10 s @20 °C; bizim sosumuzun değeri değil [A].", "Yaklaşık 1 kg ticari pizza sosu satın al; marka/lot/içindekiler ve ambalajı göster. Sulandırma veya koyulaştırma yapma.", (
        "Açtıktan sonra üstte su ayrışması var mı? Karıştırmadan önce göster.",
        "Büyük kaşıktan akışı, kesildiğinde damlama yapıp yapmadığını, 80 g dozun yayılmasını çek.",
    ), (
        ("Gerçek kayıt", "Marka/lot: ______  Sıcaklık: ____ °C  En iri parça: ____ mm"),
        ("Gözlem", "□ Akıcı  □ Yoğun  □ Ayrışıyor  □ Damlıyor  Video: __________"),
    ))

    product(doc, "2 | LAHMACUN HARCI", "110 g / pide · VARSAYIM: tasarım dozu; bir yayımlanmış deney 100 g kullanmış [B]", "Harran Üniversitesi'nin 2025 tarihli deney tarifi [B]; tek 'standart' lahmacun tarifi değildir. Antep coğrafi işaretinde soğan yok, Urfa'da var [C,D].", "Numuneyi tedarikçiden gerçek tarif olarak al. Hazır numune yoksa [B] araştırma tarifi: 500 g çiğ kıyma + 200 g soğan + 10 g isot + 120 g biber salçası + 20 g maydanoz + 25 g sarımsak + 100 mL ayçiçek yağı + 5 g karabiber + 20 g tuz. Bu, makine için onaylanmış reçete DEĞİLDİR.", (
        "İçindekileri ve gramlarını çek; kıyma çekim sayısı, doğrama boyu ve eklenen suyu yaz.",
        "Karıştırmadan önce dipte ayrılan suyu, sonra büyük kaşıktan akışı ve 110 g'ın hamura dağılımını çek.",
    ), (
        ("Gerçek kayıt", "Tarif/tedarikçi: ______  Sıcaklık: ____ °C  En iri parça: ____ mm"),
        ("Gözlem", "□ Akıyor  □ Kaşıkta kalıyor  □ Su ayrılıyor  Video: __________"),
    ))

    doc.add_page_break()
    product(doc, "3 | KUŞBAŞI ÇİĞ ET", "145 g / pide · AUTOKITCH mevcut model hedefi", "Et ve Süt Kurumu olağan kuşbaşıyı 3–4 cm olarak tanımlar [E]. Bizim 15–20 mm hedefimiz VARSAYIM ve ayrı kesim talebidir.", "Yaklaşık 1 kg çiğ sığır eti; 15–20 mm kesimli örnek iste. İmkân varsa olağan 30–40 mm kuşbaşıyı ayrı kapta yan yana göster; ikisini karıştırma.", (
        "Cetvel üzerinde en az 10 parçanın yaklaşık en/boyunu göster; en büyük parçayı ayrıca kaydet.",
        "Kabı eğerek serbest et suyunu göster; parçaların birbirine yapışması ve 145 g dökümünü çek.",
    ), (
        ("Gerçek kayıt", "Et tedarikçisi: ______  Sıcaklık: ____ °C  En iri parça: ____ mm"),
        ("Gözlem", "□ Serbest su var  □ Köprüleniyor  □ Eziliyor  Video: __________"),
    ))

    product(doc, "4 | ÇİĞ KIYMA", "160 g / pide · AUTOKITCH mevcut model hedefi", "3 mm kıyma plakası yalnızca VARSAYIM numune başlangıcıdır; gerçek ürünün parça boyu, yağ oranı ve akışkanlığı ölçülmedi.", "Yaklaşık 1 kg taze çiğ sığır kıyma; mümkünse kasaptan kullanılan plaka çapını ve çekim sayısını yazılı al. Su ekleme.", (
        "Büyük kaşıktan düşme/uzama/sürünme davranışını ve kabın dibindeki serbest sıvıyı çek.",
        "160 g porsiyonu hamura bırak; tek noktada kalıyor mu yoksa yayılıyor mu göster.",
    ), (
        ("Gerçek kayıt", "Et/yağ bilgisi: ______  Plaka: ____ mm  Sıcaklık: ____ °C"),
        ("Gözlem", "□ Yapışıyor  □ Sürüyor  □ Su ayrılıyor  Video: __________"),
    ))

    doc.add_heading("KAYNAKLAR VE SINIR", level=1)
    text(doc, "[A] Billabong Produce, Pizza Sauce Product Specification: https://www.5ways.com.au/documents/Recipes/BCT_1.pdf")
    text(doc, "[B] Aydemir ve ark. (2025), Food and Health: https://doi.org/10.3153/FH25011")
    text(doc, "[C] TÜRKPATENT Antep Lahmacunu, tescil 236: https://ci.turkpatent.gov.tr/Files/GeographicalSigns/236.pdf")
    text(doc, "[D] TÜRKPATENT Urfa Lahmacunu, tescil 353: https://ci.turkpatent.gov.tr/Files/GeographicalSigns/353.pdf")
    text(doc, "[E] Et ve Süt Kurumu, kırmızı et tanımları: https://www.esk.gov.tr/tr/10998/ET")
    text(doc, "Bu kaynaklar sadece araştırma başlangıcıdır. Onların viskozitesi, reçetesi veya parça boyu AUTOKITCH'ın gerçek satın alınacak ürününü ölçmüş sayılmaz. Nihai ekipman seçimi için aynı ürünle tedarikçi testi gerekir.")
    doc.save(VIDEO)


if __name__ == "__main__":
    supplier_form()
    video_form()
    print(SUPPLIER)
    print(VIDEO)
