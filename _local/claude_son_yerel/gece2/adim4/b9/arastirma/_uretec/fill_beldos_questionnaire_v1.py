"""Fill Beldos's original questionnaire with AUTOKITCH design targets.

No measured rheology, ingredient temperature, or operating throughput is inferred.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1] / "TEDARIK"
SOURCE = ROOT / "Beldos_Questionnaire_Original.docx"
TARGET = ROOT / "Beldos_Questionnaire_AUTOKITCH_DRAFT_v1.docx"


def put(cell, value, size=8.5):
    cell.text = value
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(size)


doc = Document(SOURCE)
header = doc.sections[0].header.tables[0]
put(header.cell(0, 1), "AUTOKITCH (project; legal entity TBC)", 9)
put(header.cell(0, 3), "Kursat Kemal Kul | k.kemalkul@gmail.com", 9)

table = doc.tables[0]

# Keep each question together when the supplier's landscape table crosses a page.
for row in table.rows:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

products = [
    (
        "Pizza sauce",
        "80 g target*",
        "1 Liquid or 2 Semi-liquid — VIDEO TO CONFIRM",
        "TBC; no measured product temperature",
        "Tomato sauce; particle size TBC",
        "Mixing batch TBC. Hopper target: 10–15 kg*",
    ),
    (
        "Lahmacun mixture",
        "110 g target*",
        "2 Semi-liquid or 3 Soft paste — VIDEO TO CONFIRM",
        "Chilled; actual °C TBC",
        "Raw mince, onion, pepper, tomato paste; largest piece TBC",
        "Mixing batch TBC. Hopper target: 2 × 21.6 kg*",
    ),
    (
        "Diced raw meat",
        "145 g target*",
        "Chunk product; viscosity scale not applicable",
        "Chilled; actual °C TBC",
        "Wet raw meat pieces, approx. 15–20 mm*; confirm on sample",
        "Mixing batch TBC. Hopper target: 5.8 kg*",
    ),
    (
        "Raw minced meat",
        "160 g target*",
        "3 Soft paste or 4 Heavy paste — VIDEO TO CONFIRM",
        "Chilled; actual °C TBC",
        "Raw minced meat; grind size TBC",
        "Mixing batch TBC. Hopper target: 6.4 kg*",
    ),
]

for row_index, values in enumerate(products, start=1):
    row = table.rows[row_index]
    row.height = None
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for col_index, value in enumerate(values):
        put(row.cells[col_index], value, 8)
    put(row.cells[6], "TBC — line not yet operating", 8)
    put(row.cells[7], "TBC — line not yet operating", 8)

# Leave the fifth product line blank: only four products were discussed with Beldos.
table.rows[5].height = None
table.rows[5].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

put(table.cell(6, 3), "Not applicable", 9)
put(table.cell(7, 3), "Not yet operating; planned schedule TBC", 9)
put(table.cell(8, 3), "Not yet operating. Design target: approx. 120 finished products/hour total*; future increase TBC.", 9)
put(table.cell(9, 3), "Onto an uncooked round pide/lahmacun base, approx. Ø280 mm* (design size; sketch/video to follow).", 9)
put(table.cell(10, 3), "To discuss. We plan separate hygienic product paths and outlets for the four products.", 9)
put(table.cell(11, 3), "Internet search", 9)

note = table.cell(12, 0).paragraphs[-1]
run = note.add_run("\nAUTOKITCH note: * engineering design targets, not measured product properties. The line is not yet operating, so no current production-process video exists. Fresh-product spoon-flow videos will follow. Please advise on suitability before selection.")
run.font.size = Pt(8)
run.bold = True

doc.save(TARGET)
print(TARGET)
